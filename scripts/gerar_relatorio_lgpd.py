from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ARQUIVO = "Relatorio_LGPD_Prato_Certo.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def force_black(style):
    style.font.color.rgb = RGBColor(0, 0, 0)
    p_pr = style.element.get_or_add_pPr()
    for border in list(p_pr.findall(qn("w:pBdr"))):
        p_pr.remove(border)


def remove_borders(paragraph):
    p_pr = paragraph._element.get_or_add_pPr()
    for border in list(p_pr.findall(qn("w:pBdr"))):
        p_pr.remove(border)


def set_cell_text(cell, text, bold=False, color="000000"):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(9.5)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_paragraph(doc, text):
    paragraph = doc.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.08
    return paragraph


def add_bullets(doc, items):
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.05
        paragraph.add_run(item)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, "1F4E3A")
        set_cell_text(cell, header, bold=True, color="FFFFFF")
        cell.width = Cm(widths[index])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_text(cells[index], value)
            cells[index].width = Cm(widths[index])
            if row_index % 2 == 1:
                set_cell_shading(cells[index], "F3F6F2")
    doc.add_paragraph()


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.9)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(18)
    styles["Title"].font.bold = True
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(13)
    styles["Heading 1"].font.bold = True
    styles["Heading 2"].font.name = "Arial"
    styles["Heading 2"].font.size = Pt(11.5)
    styles["Heading 2"].font.bold = True
    for style_name in ["Title", "Heading 1", "Heading 2"]:
        force_black(styles[style_name])

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Relatório LGPD do Aplicativo Prato Certo")
    remove_borders(title)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("Sistema de Informação  Aplicação da LGPD em projetos de tecnologia")
    run.font.size = Pt(10)
    run.bold = True

    metadata = [
        ("Projeto analisado", "Prato Certo"),
        ("Finalidade", "Aplicativo escolar para consulta de cardápio, confirmação de refeições, feedbacks e apoio à produção."),
        ("Data", "13/10/2026"),
        ("Lei considerada", "Lei Geral de Proteção de Dados LGPD"),
    ]
    add_table(doc, ["Campo", "Informação"], metadata, [4.4, 12.4])

    doc.add_heading("Objetivo do relatório", level=1)
    add_paragraph(
        doc,
        "Este relatório apresenta onde o aplicativo Prato Certo coleta, usa e armazena dados pessoais e quais medidas de LGPD foram aplicadas no projeto. O foco é mostrar como o app reduz exposição de dados, separa permissões por escola e oferece ao usuário formas claras de consentimento, consulta e solicitação de exclusão.",
    )

    doc.add_heading("Onde há coleta, uso e armazenamento de dados", level=1)
    add_paragraph(
        doc,
        "O Prato Certo usa dados pessoais para identificar o aluno, vincular a conta à escola correta e permitir o planejamento das refeições. Esses dados aparecem principalmente nos fluxos de cadastro, login, confirmação de consumo, avaliação da refeição, restrições alimentares e painel administrativo.",
    )
    add_table(
        doc,
        ["Ponto do app", "Dados tratados", "Finalidade"],
        [
            ("Cadastro e login", "Nome, matrícula, turma, escola, turnos, e-mail, senha e aceite dos documentos.", "Criar a conta, confirmar identidade básica e liberar acesso ao ambiente da escola correta."),
            ("Cardápio e confirmação", "UID do aluno, nome, turma, turno, escola, prato, data e resposta sobre consumo.", "Ajudar a equipe a planejar a quantidade de refeições sem misturar dados entre escolas."),
            ("Feedback do aluno", "Nota, emoji, etiquetas e comentário opcional vinculados à refeição e ao aluno.", "Avaliar a qualidade das refeições e apoiar melhorias no cardápio."),
            ("Restrições alimentares", "Declaração opcional de intolerâncias, alergias ou outras restrições, com responsável quando aplicável.", "Permitir planejamento mais cuidadoso para casos sensíveis, sem copiar esses dados para relatórios comuns."),
            ("Equipe administrativa", "Dados de cardápio, produção, relatórios e solicitações LGPD da própria escola.", "Gerenciar o serviço de alimentação e responder pedidos de privacidade."),
        ],
        [4.2, 6.4, 6.2],
    )

    doc.add_heading("Medidas de proteção aplicadas", level=1)
    add_bullets(
        doc,
        [
            "Termos de Uso e Política de Privacidade ficam disponíveis no app e o cadastro exige aceite obrigatório.",
            "O perfil registra a versão dos termos e a data do consentimento, permitindo comprovar qual texto foi aceito.",
            "A senha não é salva no perfil escolar; a autenticação fica no Firebase Authentication.",
            "As regras do Firestore controlam leitura e escrita por papel de conta, e-mail verificado e escola do usuário.",
            "Cada escola acessa apenas seus próprios cardápios, respostas, relatórios, restrições alimentares e solicitações LGPD.",
            "Dados sensíveis de saúde alimentar ficam em coleção separada, com consentimento específico e opção de retirada pelo aluno.",
            "Relatórios de produção trabalham com quantidades e indicadores, evitando listar nomes de alunos quando o objetivo é análise de consumo.",
            "O app registra solicitação de exclusão de dados pelo perfil do aluno, para que a escola avalie exclusão, anonimização ou conservação obrigatória.",
        ],
    )

    doc.add_heading("Como a LGPD aparece no app", level=1)
    add_paragraph(
        doc,
        "A aplicação da LGPD não ficou apenas no texto do relatório. O app recebeu telas e regras que tornam a privacidade parte do funcionamento normal do sistema. No cadastro, o aluno só consegue concluir a conta após aceitar os documentos. No perfil, há uma seção chamada Privacidade, com acesso aos documentos e um botão para pedir revisão ou exclusão de dados.",
    )
    add_paragraph(
        doc,
        "Quando o aluno envia esse pedido, o app grava um documento em solicitacoesPrivacidade usando o próprio UID. As regras impedem que ele crie solicitação para outro aluno, altere a escola, mude o status ou apague o registro. A equipe administrativa enxerga esses pedidos no painel apenas quando pertencem à sua escola. Assim, o aplicativo cria um caminho prático para exercício de direitos sem abrir acesso indevido a dados de outros usuários.",
    )
    add_paragraph(
        doc,
        "Para restrições alimentares, que podem envolver dados sensíveis de saúde, o app usa consentimento separado. O aluno pode informar depois, declarar que não possui restrições ou selecionar itens como intolerância à lactose, alergia ao leite, alergia ao ovo, restrição ao glúten e outras opções. Esses dados ficam separados dos registros comuns e podem ser retirados pelo próprio perfil.",
    )

    doc.add_heading("Ferramentas e recursos usados", level=1)
    add_table(
        doc,
        ["Recurso", "Uso no projeto"],
        [
            ("React Native e Expo", "Interface mobile e web para aluno e equipe, com telas de cadastro, perfil, cardápio, feedback e painel administrativo."),
            ("Firebase Authentication", "Login por e-mail e senha, confirmação de e-mail, recuperação de senha e reautenticação para alterações sensíveis."),
            ("Cloud Firestore", "Banco de dados com coleções separadas para usuários, cardápios, confirmações, feedbacks, produção, restrições e solicitações LGPD."),
            ("Firestore Security Rules", "Permissões por escola, papel da conta e UID, bloqueando leitura de outra escola e gravações com identidade manipulada."),
            ("Firebase Emulators", "Testes locais das regras antes de qualquer publicação no Firebase real."),
            ("Cloud Functions", "Fluxo de autorização de novo administrador sem permitir que um aluno promova a própria conta."),
        ],
        [5.0, 11.8],
    )

    doc.add_heading("Limites e cuidados necessários", level=1)
    add_paragraph(
        doc,
        "O aplicativo ajuda na proteção de dados, mas a instituição ainda precisa definir o processo real de atendimento aos pedidos LGPD. A exclusão de dados pode exigir verificação de identidade, análise de registros escolares que precisam ser mantidos e, em alguns casos, anonimização em vez de remoção imediata. Também é necessário conferir o projeto Firebase real, publicar as regras corretas e validar os fluxos com contas reais antes de uso com alunos.",
    )

    doc.add_heading("Conclusão", level=1)
    add_paragraph(
        doc,
        "O Prato Certo aplica os principais pontos pedidos no trabalho: identifica onde há coleta e armazenamento de dados, inclui termos e política de privacidade, solicita consentimento, protege dados sensíveis em coleção separada, controla acesso por escola e oferece opção de retirada ou solicitação de exclusão de dados pelo usuário. Com essas medidas, o projeto passa a tratar a privacidade como parte do funcionamento do aplicativo, e não apenas como um texto externo.",
    )
    doc.add_heading("Checklist das exigências", level=1)
    add_table(
        doc,
        ["Exigência", "Aplicação no Prato Certo"],
        [
            ("Termos de uso e política de privacidade", "Disponíveis na tela inicial, no cadastro e no perfil. O cadastro exige aceite antes de criar a conta."),
            ("Consentimento do usuário", "O perfil grava data e versão dos termos. As restrições alimentares pedem autorização específica."),
            ("Proteção de dados sensíveis", "Restrições alimentares ficam separadas, não entram em relatórios comuns e podem ser retiradas pelo aluno."),
            ("Controle de acesso", "Regras do Firestore limitam dados por UID, tipo de conta, e-mail confirmado e escola."),
            ("Exclusão de dados", "O perfil tem botão para pedido de revisão ou exclusão de dados. A equipe da escola vê o pedido e executa o processo institucional adequado."),
        ],
        [5.1, 11.7],
    )

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.add_run("Prato Certo  Relatório LGPD").font.size = Pt(8)

    doc.save(ARQUIVO)


if __name__ == "__main__":
    main()
